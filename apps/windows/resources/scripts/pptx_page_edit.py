#!/usr/bin/env python3
"""PPT 单页文本重写（P2 T4）

只改目标页的标题占位符与「最大的正文文本框」；其余页与其它 shape 不动。
由主进程用受管 Python + python-pptx 调用，参数走 argv（不经 shell）。

用法:
  python pptx_page_edit.py --input in.pptx --output out.pptx --slide 2 \
      --spec '{"title":"...","bullets":["...","..."]}'

输出（stdout 一行 JSON）:
  {"ok": true, "slideCount": 4, "changed": {"title": true, "bullets": 2}}
  {"ok": false, "error": "..."}  (非零退出)
"""

import argparse
import json
import sys
from copy import deepcopy

try:
    from pptx import Presentation
except Exception as exc:  # pragma: no cover - dependency error path
    print(json.dumps({"ok": False, "error": f"python-pptx 不可用: {exc}"}, ensure_ascii=False))
    sys.exit(1)


MAX_BULLETS = 8
MAX_CHARS = 500


def fail(message: str) -> None:
    print(json.dumps({"ok": False, "error": message}, ensure_ascii=False))
    sys.exit(1)


def set_para_text(paragraph, text: str) -> None:
    """保留首段/首个 run 的字体与段落属性，只替换文字。"""
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in list(paragraph.runs)[1:]:
            run._r.getparent().remove(run._r)
    else:
        paragraph.text = text


def remove_trailing_paragraphs(text_frame, keep: int) -> None:
    while len(text_frame.paragraphs) > keep:
        p = text_frame.paragraphs[-1]._p
        p.getparent().remove(p)


def body_shape(slide):
    title = slide.shapes.title
    candidates = [
        shape
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
        and shape.has_text_frame
        and shape is not title
        and (shape.width or 0) * (shape.height or 0) > 0
    ]
    if not candidates:
        return None
    return max(candidates, key=lambda s: (s.width or 0) * (s.height or 0))


def rewrite_title(slide, title: str) -> None:
    shape = slide.shapes.title
    if shape is None or not getattr(shape, "has_text_frame", False):
        raise ValueError("这一页没有标题占位符")
    tf = shape.text_frame
    if not tf.paragraphs:
        raise ValueError("标题为空")
    set_para_text(tf.paragraphs[0], title)
    remove_trailing_paragraphs(tf, 1)


def rewrite_bullets(slide, bullets) -> None:
    shape = body_shape(slide)
    if shape is None:
        raise ValueError("这一页没有可编辑的正文文本框")
    tf = shape.text_frame
    if not tf.paragraphs:
        tf.paragraphs[0].text = bullets[0]
    else:
        first = tf.paragraphs[0]
        set_para_text(first, bullets[0])
        template = deepcopy(first._p)
        remove_trailing_paragraphs(tf, 1)
        for text in bullets[1:]:
            tf._txBody.append(deepcopy(template))
            set_para_text(tf.paragraphs[-1], text)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--slide", required=True, type=int)
    parser.add_argument("--spec", required=True)
    args = parser.parse_args()

    try:
        spec = json.loads(args.spec)
    except Exception as exc:
        fail(f"spec 不是合法 JSON: {exc}")

    title = spec.get("title")
    bullets = spec.get("bullets")
    if title is not None:
        if not isinstance(title, str) or not title.strip():
            fail("title 必须是非空字符串")
        title = title.strip()[:MAX_CHARS]
    if bullets is not None:
        if not isinstance(bullets, list) or not bullets:
            fail("bullets 必须是非空数组")
        if len(bullets) > MAX_BULLETS:
            fail(f"bullets 最多 {MAX_BULLETS} 条")
        cleaned = []
        for item in bullets:
            if not isinstance(item, str) or not item.strip():
                fail("每条 bullet 必须是非空字符串")
            cleaned.append(item.strip()[:MAX_CHARS])
        bullets = cleaned
    if title is None and bullets is None:
        fail("spec 里没有 title 或 bullets")

    try:
        prs = Presentation(args.input)
        if args.slide < 0 or args.slide >= len(prs.slides):
            fail(f"页码越界: {args.slide}（共 {len(prs.slides)} 页）")
        slide = prs.slides[args.slide]
        if title is not None:
            rewrite_title(slide, title)
        if bullets is not None:
            rewrite_bullets(slide, bullets)
        prs.save(args.output)
        print(
            json.dumps(
                {
                    "ok": True,
                    "slideCount": len(prs.slides),
                    "changed": {
                        "title": title is not None,
                        "bullets": len(bullets) if bullets is not None else 0,
                    },
                },
                ensure_ascii=False,
            )
        )
    except SystemExit:
        raise
    except Exception as exc:
        fail(str(exc))


if __name__ == "__main__":
    main()
